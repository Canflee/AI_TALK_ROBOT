import logging
import os
import re
from pathlib import Path

from ..config import BASE_DIR

logger = logging.getLogger(__name__)

PROCESSED_DIR = BASE_DIR / "data" / "processed"
FONT_CANDIDATES = (
    Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "simhei.ttf",
    Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "Deng.ttf",
)

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^[-*•]\s+(.*)$")
_NUMBERED = re.compile(r"^\d+[.)、]\s+(.*)$")


def write_outputs(
    text: str,
    job_id: int,
    source_name: str,
    formats: tuple[str, ...] = ("docx", "pdf"),
) -> tuple[Path | None, Path | None]:
    folder = PROCESSED_DIR / str(job_id)
    folder.mkdir(parents=True, exist_ok=True)
    stem = _safe_stem(Path(source_name).stem) or "处理结果"
    docx_path = folder / f"{stem}_处理结果.docx" if "docx" in formats else None
    pdf_path = folder / f"{stem}_处理结果.pdf" if "pdf" in formats else None

    if docx_path is not None:
        write_docx(text, docx_path, stem)
    if pdf_path is not None:
        try:
            write_pdf(text, pdf_path, stem)
        except Exception:
            logger.exception("生成 PDF 失败: %s", pdf_path)
            pdf_path = None
    return docx_path, pdf_path


def write_docx(text: str, path: Path, title: str) -> None:
    from docx import Document
    from docx.oxml.ns import qn

    document = Document()
    _set_docx_font(document.styles["Normal"], "Microsoft YaHei", "微软雅黑")
    document.add_heading(title, level=0)

    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        heading = _HEADING.match(stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            document.add_heading(heading.group(2).strip(), level=level)
            continue
        bullet = _BULLET.match(stripped)
        if bullet:
            paragraph = document.add_paragraph(bullet.group(1).strip(), style="List Bullet")
        elif _NUMBERED.match(stripped):
            paragraph = document.add_paragraph(
                _NUMBERED.match(stripped).group(1).strip(), style="List Number"
            )
        else:
            paragraph = document.add_paragraph(stripped)
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def write_pdf(text: str, path: Path, title: str) -> None:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        from xml.sax.saxutils import escape
    except ImportError as exc:
        raise RuntimeError("缺少 reportlab，无法生成 PDF") from exc

    font_name = _register_pdf_font(pdfmetrics, TTFont)
    body_style = ParagraphStyle(
        "DocBody",
        fontName=font_name,
        fontSize=11,
        leading=18,
        wordWrap="CJK",
        spaceAfter=6,
    )
    heading_style = ParagraphStyle(
        "DocHeading",
        parent=body_style,
        fontSize=15,
        leading=22,
        spaceBefore=10,
        spaceAfter=6,
    )
    title_style = ParagraphStyle(
        "DocTitle",
        parent=body_style,
        fontSize=20,
        leading=28,
        spaceAfter=14,
    )

    story = [Paragraph(escape(title), title_style), Spacer(1, 6)]
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        heading = _HEADING.match(stripped)
        if heading:
            story.append(Paragraph(escape(heading.group(2).strip()), heading_style))
            continue
        bullet = _BULLET.match(stripped)
        if bullet:
            story.append(Paragraph(f"• {escape(bullet.group(1).strip())}", body_style))
            continue
        numbered = _NUMBERED.match(stripped)
        if numbered:
            story.append(Paragraph(f"• {escape(numbered.group(1).strip())}", body_style))
            continue
        story.append(Paragraph(escape(stripped), body_style))

    path.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=title,
    )
    document.build(story)


def _register_pdf_font(pdfmetrics, TTFont) -> str:
    for candidate in FONT_CANDIDATES:
        if not candidate.exists():
            continue
        try:
            if "DocChinese" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("DocChinese", str(candidate)))
            return "DocChinese"
        except Exception:
            logger.warning("注册 PDF 中文字体失败: %s", candidate, exc_info=True)
    logger.warning("未找到可用的中文字体，PDF 中文可能显示异常")
    return "Helvetica"


def _set_docx_font(style, name: str, east_asia: str) -> None:
    from docx.oxml.ns import qn

    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)


def _safe_stem(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name or "").strip(" .")
    return name[:80]
